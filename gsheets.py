"""
Google-connected renderer (production).

Does what Code.gs does, from the server: writes the student's marks into the
Marksheet / Transcript template sheets of the live Google Sheet, then asks
Google to export those sheets to PDF. Because the templates are read from the
spreadsheet on every lookup, new or edited template sheets are used straight
away, and the PDFs look exactly like the Apps Script ones.

The service account needs Editor access to the spreadsheet (it writes the
template sheets, as the Apps Script did).
"""

import concurrent.futures
import datetime
import re
import threading
import urllib.parse

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

import engine as E

SCOPES = ["https://www.googleapis.com/auth/spreadsheets",
          "https://www.googleapis.com/auth/drive.readonly"]

PCT = {"type": "PERCENT", "pattern": "0%"}
TEXT = {"type": "TEXT"}
DATE = {"type": "DATE", "pattern": "d-mmm-yy"}


# ------------------------------- cell helpers -------------------------------

def _rc(a1):
    """'E8' -> (8, 5)"""
    col, row = coordinate_from_string(a1)
    return row, column_index_from_string(col)


def _isnum(s):
    try:
        float(s)
        return True
    except (TypeError, ValueError):
        return False


class Formula(str):
    """A value to be written as a formula."""


# ------------------------------- pure planning -------------------------------
#
# `cells` describes what is currently in the template: {(row, col): Cell} for
# the cells read below, where Cell has .text (shown value), .formula (bool) and
# .is_date (bool). The planners return a list of writes (row, col, value, fmt):
# value is a number, a str ('' clears the cell) or a Formula; fmt is a
# numberFormat dict or None (leave the format alone).

class Cell:
    def __init__(self, text="", formula=False, is_date=False):
        self.text, self.formula, self.is_date = text, formula, is_date


def _module_count(M):
    last = column_index_from_string(re.search(r":([A-Z]+)\d+$", M["range"]).group(1))
    return last - M["firstCol"] + 1


def marksheet_reads(program):
    M = program["marksheet"]
    cells = [(M["moduleNameRow"], M["firstCol"] + i) for i in range(_module_count(M))]
    cells += [_rc(M["startCell"]), _rc(M["endCell"])]
    if M.get("finalMarkCell"):
        cells.append(_rc(M["finalMarkCell"]))
    return cells


def transcript_reads(T):
    cells = [(r, T["moduleCol"]) for r in range(T["firstRow"], T["lastRow"] + 1)]
    return cells + [_rc(T["startCell"]), _rc(T["endCell"]), _rc(T["averageCell"]), _rc(T["programCell"])]


def _date_write(cells, a1, value):
    r, c = _rc(a1)
    if isinstance(value, datetime.datetime):      # -> Sheets serial number
        value = (value - datetime.datetime(1899, 12, 30)).total_seconds() / 86400
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # Keep the template's own date format; give a plain cell a date format
        return (r, c, value, None if cells[(r, c)].is_date else DATE)
    return (r, c, "" if value is None else value, None)


def plan_marksheet(program, student, cells):
    """Mirror of fillMarksheet_ in Code.gs."""
    M = program["marksheet"]
    names = [cells[(M["moduleNameRow"], M["firstCol"] + i)].text for i in range(_module_count(M))]

    b = E.build_marksheet_rows(student, names, program)
    out = []
    for r in b["rows"]:
        for i, v in enumerate(r["values"]):
            if v is not None:
                out.append((r["row"], M["firstCol"] + i, v, PCT))

    sid = student["id"]
    out.append(_rc(M["idCell"]) + ((float(sid) if _isnum(sid) else sid), None))
    out.append(_rc(M["nameCell"]) + (student["fullName"], None))
    out.append(_date_write(cells, M["startCell"], student["start"]))
    out.append(_date_write(cells, M["endCell"], student["finish"]))
    if M.get("finalMarkCell") and not cells[_rc(M["finalMarkCell"])].formula:
        out.append(_rc(M["finalMarkCell"]) + (b["finalMark"], None))
    return out


def plan_transcript(program, student, T, cells):
    """Mirror of fillTranscript_ in Code.gs."""
    count = T["lastRow"] - T["firstRow"] + 1
    names = [cells[(T["firstRow"] + i, T["moduleCol"])].text for i in range(count)]

    t = E.build_transcript_block(student, names, program)
    out = []
    for i, e in enumerate(t["rows"]):
        row = T["firstRow"] + i
        if e["kind"] in ("mark", "empty"):
            out.append((row, T["earnedCol"], e["earned"], None))
            out.append((row, T["gradeCol"], e["grade"], None))
        elif e["kind"] == "status":
            col = (T.get("statusCols") or {}).get(e["key"]) or T.get("statusCol") or 8
            out.append((row, col, e["text"], None))

    if not cells[_rc(T["averageCell"])].formula:
        out.append(_rc(T["averageCell"]) + (t["average"], None))
    out.append(_rc(T["averageGradeCell"]) + (t["averageGrade"], None))
    out.append(_rc(T["nameCell"]) + (student["fullName"], None))
    out.append(_rc(T["idCell"]) + (str(student["id"]), TEXT))
    out.append(_date_write(cells, T["startCell"], student["start"]))
    out.append(_date_write(cells, T["endCell"], student["finish"]))

    if T.get("credentialCell"):
        col = get_column_letter(T["earnedCol"])
        marks = "%s%d:%s%d" % (col, T["firstRow"], col, T.get("marksLastRow") or T["lastRow"])
        out.append(_rc(T["credentialCell"]) + (Formula(E.credential_formula_(marks, T["passMark"])), None))
    return out


# ------------------------------- Google API -------------------------------

def _a1(sheet, row, col):
    return "'%s'!%s%d" % (sheet.replace("'", "''"), get_column_letter(col), row)


def _update_request(sheet_id, row, col, value, fmt):
    if isinstance(value, Formula):
        cell = {"userEnteredValue": {"formulaValue": str(value)}}
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        cell = {"userEnteredValue": {"numberValue": value}}
    elif value == "":
        cell = {}                                  # clears the value
    else:
        cell = {"userEnteredValue": {"stringValue": str(value)}}
    fields = "userEnteredValue"
    if fmt:
        cell["userEnteredFormat"] = {"numberFormat": fmt}
        fields += ",userEnteredFormat.numberFormat"
    return {"updateCells": {
        "rows": [{"values": [cell]}],
        "fields": fields,
        "start": {"sheetId": sheet_id, "rowIndex": row - 1, "columnIndex": col - 1},
    }}


class GoogleSheetRenderer:
    def __init__(self, spreadsheet_id, sa_json_path):
        from google.oauth2.service_account import Credentials
        from google.auth.transport.requests import AuthorizedSession
        from googleapiclient.discovery import build
        creds = Credentials.from_service_account_file(sa_json_path, scopes=SCOPES)
        self._svc = build("sheets", "v4", credentials=creds, cache_discovery=False)
        self._http = AuthorizedSession(creds)
        self._id = spreadsheet_id
        # The templates are shared sheets, so only one lookup may use them at a time
        self._lock = threading.Lock()

    def _read(self, wanted):
        """wanted: {sheet_name: [(row, col), ...]} -> ({sheet: sheetId}, {sheet: {(row, col): Cell}})"""
        from googleapiclient.errors import HttpError
        ranges, order = [], []
        for sheet, rcs in wanted.items():
            for rc in rcs:
                ranges.append(_a1(sheet, *rc))
                order.append((sheet, rc))
        try:
            resp = self._svc.spreadsheets().get(
                spreadsheetId=self._id, ranges=ranges,
                fields="sheets(properties(sheetId,title),data(rowData(values("
                       "userEnteredValue,formattedValue,effectiveFormat/numberFormat))))",
            ).execute()
        except HttpError as e:
            if e.resp.status == 400:
                raise ValueError('Sheet %s was not found.' % " / ".join('"%s"' % s for s in wanted))
            raise

        ids, cells = {}, {sheet: {} for sheet in wanted}
        by_sheet = {}
        for s in resp.get("sheets", []):
            ids[s["properties"]["title"]] = s["properties"]["sheetId"]
            by_sheet[s["properties"]["title"]] = s.get("data", [])
        pos = {sheet: 0 for sheet in wanted}
        for sheet, rc in order:               # data blocks come back in request order per sheet
            block = by_sheet[sheet][pos[sheet]]
            pos[sheet] += 1
            v = ((block.get("rowData") or [{}])[0].get("values") or [{}])[0]
            fmt_type = (v.get("effectiveFormat", {}).get("numberFormat") or {}).get("type", "")
            cells[sheet][rc] = Cell(text=v.get("formattedValue", ""),
                                    formula="formulaValue" in v.get("userEnteredValue", {}),
                                    is_date=fmt_type in ("DATE", "DATE_TIME"))
        return ids, cells

    def _export(self, gid, a1range):
        P = E.PDF_SETTINGS
        params = {
            "format": "pdf", "gid": gid, "range": a1range, "size": "letter",
            "portrait": str(P["portrait"]).lower(), "scale": 4,
            "top_margin": P["margin"], "bottom_margin": P["margin"],
            "left_margin": P["margin"], "right_margin": P["margin"],
            "horizontal_alignment": "CENTER", "vertical_alignment": "TOP",
            "gridlines": "false", "printtitle": "false", "sheetnames": "false",
            "pagenumbers": "false", "fzr": "false", "fzc": "false",
        }
        url = "https://docs.google.com/spreadsheets/d/%s/export?%s" % (
            self._id, urllib.parse.urlencode(params))
        res = self._http.get(url, timeout=120)
        if res.status_code != 200:
            raise RuntimeError("PDF export failed (HTTP %d). Check that the spreadsheet is shared "
                               "with the service account as Editor." % res.status_code)
        return res.content

    def generate(self, source, raw_id, program_key=None):
        """Same result as getStudentDocuments in Code.gs, with the PDFs as bytes."""
        found = E.lookup(source, raw_id, program_key)
        if found.get("ambiguous"):
            return found
        program, student = found["program"], found["student"]
        tr_sheet, T = E.choose_transcript(program, student)
        ms_sheet = program["marksheetSheet"]

        with self._lock:
            ids, cells = self._read({ms_sheet: marksheet_reads(program), tr_sheet: transcript_reads(T)})
            requests = [_update_request(ids[ms_sheet], *w) for w in plan_marksheet(program, student, cells[ms_sheet])]
            requests += [_update_request(ids[tr_sheet], *w) for w in plan_transcript(program, student, T, cells[tr_sheet])]
            self._svc.spreadsheets().batchUpdate(spreadsheetId=self._id, body={"requests": requests}).execute()

            with concurrent.futures.ThreadPoolExecutor(2) as pool:
                ms_pdf = pool.submit(self._export, ids[ms_sheet], program["marksheet"]["range"])
                tr_pdf = pool.submit(self._export, ids[tr_sheet], T["range"])
                ms_pdf, tr_pdf = ms_pdf.result(), tr_pdf.result()

        safe = re.sub(r'[\\/:*?"<>|]+', "", student["fullName"]).strip()
        return {
            "studentId": student["id"],
            "studentName": student["fullName"],
            "program": cells[tr_sheet][_rc(T["programCell"])].text.strip() or program["label"],
            "filename": "%s - %s - Marksheet & Transcript.pdf" % (student["id"], safe),
            "marksheetPdf": ms_pdf,
            "transcriptPdf": tr_pdf,
        }
