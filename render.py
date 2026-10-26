"""
Template filling + PDF rendering.

Takes the engine's computed values, writes them into the .xlsx template sheets
(openpyxl, preserving all existing formatting/images/formulas), then renders
each sheet to PDF with LibreOffice headless and merges the two into one file.

This is the replacement for the Apps Script fillMarksheet_/fillTranscript_/
exportPdf_ functions. The Google Sheets export engine is gone; LibreOffice
renders the same sheets instead.
"""

import os
import re
import shutil
import subprocess
import tempfile
import datetime

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.page import PageMargins
from pypdf import PdfReader, PdfWriter

import engine as E

SOFFICE = shutil.which("soffice") or shutil.which("libreoffice") or "soffice"


# ------------------------------- cell helpers -------------------------------

def _is_formula(cell):
    v = cell.value
    return isinstance(v, str) and v.startswith("=")


def _set_unless_formula(ws, a1, value):
    if not _is_formula(ws[a1]):
        ws[a1] = value


def _range_last_col(a1range):
    m = re.search(r":([A-Z]+)\d+$", a1range)
    return column_index_from_string(m.group(1))


# ------------------------------- filling -------------------------------

def fill_marksheet(ws, program, student):
    M = program["marksheet"]
    last_col = _range_last_col(M["range"])
    n = last_col - M["firstCol"] + 1
    names = [ws.cell(M["moduleNameRow"], M["firstCol"] + i).value for i in range(n)]

    b = E.build_marksheet_rows(student, names, program)
    for r in b["rows"]:
        for i, v in enumerate(r["values"]):
            if v is None:
                continue  # leave the cell alone
            cell = ws.cell(r["row"], M["firstCol"] + i)
            cell.number_format = "0%"
            cell.value = v

    sid = student["id"]
    ws[M["idCell"]] = (float(sid) if _isnum(sid) else sid)
    ws[M["nameCell"]] = student["fullName"]
    ws[M["startCell"]] = _as_date(student["start"])
    ws[M["endCell"]] = _as_date(student["finish"])
    if M.get("finalMarkCell"):
        _set_unless_formula(ws, M["finalMarkCell"], b["finalMark"])


def fill_transcript(ws, program, student, T):
    count = T["lastRow"] - T["firstRow"] + 1
    names = [ws.cell(T["firstRow"] + i, T["moduleCol"]).value for i in range(count)]

    t = E.build_transcript_block(student, names, program)
    rows = t["rows"]

    i = 0
    while i < count:
        e = rows[i]
        if e["kind"] in ("mark", "empty"):
            while i < count and rows[i]["kind"] in ("mark", "empty"):
                ws.cell(T["firstRow"] + i, T["earnedCol"]).value = rows[i]["earned"]
                ws.cell(T["firstRow"] + i, T["gradeCol"]).value = rows[i]["grade"]
                i += 1
        else:
            if e["kind"] == "status":
                col = (T.get("statusCols") or {}).get(e["key"]) or T.get("statusCol") or 8
                ws.cell(T["firstRow"] + i, col).value = e["text"]
            i += 1

    _set_unless_formula(ws, T["averageCell"], t["average"])
    ws[T["averageGradeCell"]] = t["averageGrade"]
    ws[T["nameCell"]] = student["fullName"]
    idc = ws[T["idCell"]]
    idc.number_format = "@"
    idc.value = str(student["id"])
    ws[T["startCell"]] = _as_date(student["start"])
    ws[T["endCell"]] = _as_date(student["finish"])

    if T.get("credentialCell"):
        last_mark_row = T.get("marksLastRow") or T["lastRow"]
        col = get_column_letter(T["earnedCol"])
        marks = "%s%d:%s%d" % (col, T["firstRow"], col, last_mark_row)
        ws[T["credentialCell"]] = E.credential_formula_(marks, T["passMark"])


def _as_date(v):
    """Sheets API serial number (days since 1899-12-30) -> datetime; anything else as-is."""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=v)
    return v


def _isnum(s):
    try:
        float(s)
        return True
    except (TypeError, ValueError):
        return False


# ------------------------------- page setup + PDF -------------------------------

def _apply_page_setup(ws, a1range):
    # Honour the template's own orientation (the marksheet is authored landscape,
    # the transcript portrait). Margins and alignment follow the Apps Script export:
    # 0.25" margins, centered across, pinned to the top.
    m = E.PDF_SETTINGS["margin"]
    ws.page_margins = PageMargins(left=m, right=m, top=m, bottom=m, header=0, footer=0)
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = False
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    pr = ws.sheet_properties.pageSetUpPr
    if pr is None:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    else:
        pr.fitToPage = True
    ws.print_area = a1range


# One warm LibreOffice profile reused for the life of the process. The cold
# first-run init is the slow part; paying it once makes each conversion quick.
_LO_PROFILE = os.path.join(tempfile.gettempdir(), "mt_lo_profile")


def _prepare_sheet(filled_xlsx, sheet_name, a1range, workdir, tag):
    """Write a single-sheet xlsx with print layout applied; no conversion yet."""
    wb = openpyxl.load_workbook(filled_xlsx)
    for nm in list(wb.sheetnames):
        if nm != sheet_name:
            del wb[nm]
    _apply_page_setup(wb[sheet_name], a1range)
    path = os.path.join(workdir, "doc_%s.xlsx" % tag)
    wb.save(path)
    return path


def _convert_to_pdf(xlsx_paths, workdir):
    """Convert several xlsx to PDF in one LibreOffice call. Returns pdf paths."""
    subprocess.run(
        [SOFFICE, "--headless", "--nologo", "--nofirststartwizard",
         "-env:UserInstallation=file://%s" % _LO_PROFILE,
         "--convert-to", "pdf:calc_pdf_Export", "--outdir", workdir, *xlsx_paths],
        check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180,
    )
    pdfs = []
    for x in xlsx_paths:
        p = os.path.join(workdir, os.path.splitext(os.path.basename(x))[0] + ".pdf")
        if not os.path.exists(p):
            raise RuntimeError("LibreOffice did not produce a PDF for %s" % x)
        pdfs.append(p)
    return pdfs


def _merge(pdfs, dest):
    w = PdfWriter()
    for p in pdfs:
        for pg in PdfReader(p).pages:
            w.add_page(pg)
    with open(dest, "wb") as f:
        w.write(f)


# ------------------------------- public entry -------------------------------

def generate(template_xlsx, source, raw_id, program_key=None, out_dir=None):
    """
    source: object with .values(sheet_name) -> 2D list.
    Returns dict: either {"ambiguous": [...]} or the student result with a
    combined PDF path. Mirrors getStudentDocuments in Code.gs.
    """
    sid = str("" if raw_id is None else raw_id).strip()
    if not sid:
        raise ValueError("Please enter a Student ID.")

    matches = []
    for p in E.PROGRAMS:
        if program_key and p["key"] != program_key:
            continue
        student = E.read_student(source.values(p["dataSheet"]), sid, p)
        if student:
            matches.append((p, student))

    if not matches:
        raise ValueError('No student found with ID "%s".' % sid)
    if len(matches) > 1:
        return {"ambiguous": [{"key": p["key"], "label": p["label"],
                               "name": s["fullName"]} for p, s in matches]}

    program, student = matches[0]
    pass_sheet, T = E.choose_transcript(program, student)

    work = out_dir or tempfile.mkdtemp(prefix="mt_")
    os.makedirs(work, exist_ok=True)

    # Fill a private copy of the template workbook for this request. Keep only
    # the two sheets we render, so the heavy data sheets never reach LibreOffice.
    filled = os.path.join(work, "filled.xlsx")
    wb = openpyxl.load_workbook(template_xlsx)
    keep = {program["marksheetSheet"], pass_sheet}
    for nm in list(wb.sheetnames):
        if nm not in keep:
            del wb[nm]
    fill_marksheet(wb[program["marksheetSheet"]], program, student)
    fill_transcript(wb[pass_sheet], program, student, T)
    program_text = str(wb[pass_sheet][T["programCell"]].value or "").strip()
    wb.save(filled)

    ms_xlsx = _prepare_sheet(filled, program["marksheetSheet"], program["marksheet"]["range"], work, "ms")
    tr_xlsx = _prepare_sheet(filled, pass_sheet, T["range"], work, "tr")
    ms_pdf, tr_pdf = _convert_to_pdf([ms_xlsx, tr_xlsx], work)

    safe = re.sub(r'[\\/:*?"<>|]+', "", student["fullName"]).strip()
    fname = "%s - %s - Marksheet & Transcript.pdf" % (student["id"], safe)
    combined = os.path.join(work, fname)
    _merge([ms_pdf, tr_pdf], combined)

    return {
        "studentId": student["id"],
        "studentName": student["fullName"],
        "program": program_text or program["label"],
        "filename": fname,
        "marksheetPdf": ms_pdf,
        "transcriptPdf": tr_pdf,
        "combinedPdf": combined,
    }
